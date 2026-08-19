from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from practice.site_selection import (
    ReportGenerationBlockedError,
    generate_site_selection_report,
    run_parallel_site_selection_workflow,
)
from tests.test_site_scoring import composite_dependencies
from tests.test_site_selection_workflow import manifests, request


def completed_state():
    return run_parallel_site_selection_workflow(
        request(),
        manifests(),
        composite_dependencies(),
    )


def test_report_contains_scores_policy_lineage_and_safety_boundary(
    tmp_path: Path,
) -> None:
    output = tmp_path / "site-selection-report.docx"

    generate_site_selection_report(completed_state(), output)

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    combined = text + "\n" + table_text
    assert "选址预审证据报告" in combined
    assert "REQ-workflow" in combined
    assert "site-score-fixture-1.0" in combined
    assert "第一条" in combined
    assert "policy://demo/workflow" in combined
    assert "POI 来源与时间" in combined
    assert "poi-mock" in combined
    assert "2026-08-18T23:55:00+00:00" in combined
    assert "不构成整体合规结论" in combined
    assert "未生成" in combined
    assert len(document.tables) >= 7


def test_report_tables_have_fixed_matching_dxa_geometry(tmp_path: Path) -> None:
    output = tmp_path / "site-selection-report.docx"
    generate_site_selection_report(completed_state(), output)
    document = Document(output)

    for table in document.tables:
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        table_indent = table._tbl.tblPr.find(qn("w:tblInd"))
        grid_widths = [
            int(column.get(qn("w:w")))
            for column in table._tbl.tblGrid
        ]
        assert table_width.get(qn("w:w")) == "9360"
        assert table_indent.get(qn("w:w")) == "120"
        assert sum(grid_widths) == 9360
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        for row in table.rows:
            cell_widths = [
                int(cell._tc.tcPr.tcW.get(qn("w:w")))
                for cell in row.cells
            ]
            assert cell_widths == grid_widths


def test_report_rejects_state_without_review(tmp_path: Path) -> None:
    state = completed_state()
    data = state.model_dump()
    data["evidence_review_report"] = None
    unreviewed = type(state).model_validate(data)

    with pytest.raises(ReportGenerationBlockedError, match="审查报告"):
        generate_site_selection_report(
            unreviewed,
            tmp_path / "must-not-exist.docx",
        )
