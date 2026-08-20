from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .evidence import AgentState, AnalysisResult, AnalysisStatus
class ReportGenerationBlockedError(RuntimeError):
    """Raised when an auditable report cannot be generated from the state."""


_CONTENT_WIDTH_DXA = 9360
_TABLE_INDENT_DXA = 120
_CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
_BLUE = RGBColor(0x2E, 0x74, 0xB5)
_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
_GRAY = RGBColor(0x55, 0x55, 0x55)


def generate_site_selection_report(
    state: AgentState,
    output_path: str | Path,
) -> Path:
    """Generate a decision-memo style evidence report from reviewed state."""

    _validate_report_state(state)
    document = Document()
    _configure_document(document)
    _add_title_block(document, state)
    _add_disclaimer(document)
    _add_summary(document, state)
    _add_comparison(document, state)
    for result in state.results:
        _add_candidate_section(document, result)
    _add_review_section(document, state)
    _add_source_section(document, state)

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _validate_report_state(state: AgentState) -> None:
    if state.status is not AnalysisStatus.COMPLETED:
        raise ReportGenerationBlockedError("Word 报告要求分析状态为 completed")
    if not state.results or state.comparison_report is None:
        raise ReportGenerationBlockedError("Word 报告要求完整结果和候选对比")
    if state.evidence_review_report is None:
        raise ReportGenerationBlockedError("Word 报告要求证据审查报告")


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, "Arial", 11, RGBColor(0, 0, 0))
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, _BLUE, 12, 6),
        "Heading 2": (13, _BLUE, 10, 5),
        "Heading 3": (12, _DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        _set_style_font(style, "Arial", size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    _add_run(header, "SITE SELECTION / EVIDENCE REVIEW", 8.5, _GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    _add_run(footer, "Auditable pre-screening report", 8.5, _GRAY)


def _add_title_block(document: DocumentType, state: AgentState) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    _add_run(title, "选址预审证据报告", 23, RGBColor(0, 0, 0), bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    _add_run(
        subtitle,
        f"{state.request.project_type.value} / {state.request.request_id}",
        14,
        _GRAY,
    )
    metadata = [
        ("请求时间", state.request.requested_at.isoformat()),
        ("候选地块", str(len(state.request.candidate_parcels))),
        ("分析状态", state.status.value),
        ("证据审查", state.evidence_review_report.status.value),
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        _add_run(paragraph, f"{label}: ", 11, RGBColor(0, 0, 0), bold=True)
        _add_run(paragraph, value, 11, RGBColor(0, 0, 0))


def _add_disclaimer(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(10)
    _shade_paragraph(paragraph, "F4F6F9")
    _add_run(paragraph, "使用边界: ", 10.5, _DARK_BLUE, bold=True)
    _add_run(
        paragraph,
        "本报告用于展示数据、规则命中、软评分和证据血缘，不构成整体合规结论、行政审批意见或选址推荐。未命中当前规则不等于整体合规。",
        10.5,
        RGBColor(0, 0, 0),
    )


def _add_summary(document: DocumentType, state: AgentState) -> None:
    document.add_heading("1. 审查摘要", level=1)
    report = state.evidence_review_report
    rows = [
        ("Review 版本", report.review_version),
        ("Review 状态", report.status.value),
        ("需人工复核", "是" if report.requires_human_review else "否"),
        ("评分版本", state.comparison_report.scoring_version),
        ("排名口径", state.comparison_report.ranking_basis),
    ]
    _add_label_detail_table(document, rows)


def _add_comparison(document: DocumentType, state: AgentState) -> None:
    document.add_heading("2. 候选地块对比", level=1)
    headers = ["地块", "软分名次", "软评分", "政策结果"]
    rows = []
    for item in state.comparison_report.candidates:
        outcomes = ", ".join(outcome.value for outcome in item.policy_outcomes)
        rows.append(
            [
                item.parcel_id,
                str(item.soft_rank),
                f"{item.soft_score:.2f}",
                outcomes or "未命中当前规则",
            ]
        )
    _add_table(document, headers, rows, [1800, 1800, 1800, 3960])
    _add_table_note(
        document,
        "说明：软评分名次与政策规则结果分开展示，排名不覆盖硬规则或人工复核。",
    )


def _add_candidate_section(
    document: DocumentType,
    result: AnalysisResult,
) -> None:
    document.add_heading(f"3. 候选地块 {result.parcel_id}", level=1)
    rows = [
        ("项目类型", result.project_type.value),
        ("软评分", _format_optional_score(result.overall_soft_score)),
        (
            "评分口径",
            (
                result.site_score_report.scoring_version
                if result.site_score_report is not None
                else result.poi_evidence.score_report.scoring_version
            ),
        ),
        ("自动结论", result.conclusion or "未生成"),
    ]
    _add_label_detail_table(document, rows)

    document.add_heading("3.1 GIS 证据", level=2)
    gis_rows = [
        [metric, f"{value:.6g}"]
        for metric, value in sorted(result.gis_evidence.metrics.items())
    ]
    _add_table(document, ["指标", "数值"], gis_rows, [2700, 6660])
    _add_table_note(
        document,
        "数据集: " + ", ".join(result.gis_evidence.dataset_ids)
        + f" | CRS: {result.gis_evidence.crs}",
    )

    document.add_heading("3.2 POI 评分证据", level=2)
    poi_report = result.poi_evidence.score_report
    poi_rows = [
        [
            group.group_key,
            group.source_dataset_id,
            f"{group.group_score:.2f}",
            f"{group.weighted_score:.2f}",
        ]
        for group in poi_report.group_scores
    ]
    _add_table(
        document,
        ["分组", "来源数据集", "组分", "加权分"],
        poi_rows,
        [2100, 3300, 1800, 2160],
    )
    _add_table_note(
        document,
        f"POI 评分版本: {poi_report.scoring_version} | 总分: {poi_report.total_score:.2f}",
    )

    document.add_heading("3.2.1 POI 来源与时间", level=3)
    source_rows = []
    for feature_set in result.poi_evidence.feature_sets:
        source = feature_set.source
        fallback = (
            f"{source.fallback_from.value}: {source.fallback_reason}"
            if source.fallback_from is not None
            else "无"
        )
        source_rows.append(
            [
                feature_set.query.group_key,
                source.provider.value,
                source.dataset_id,
                source.dataset_version or "未标注",
                source.queried_at.isoformat(),
                (
                    source.dataset_updated_at.isoformat()
                    if source.dataset_updated_at is not None
                    else "未标注"
                ),
                source.crs,
                (
                    f"{source.record_count}/"
                    f"{source.available_record_count if source.available_record_count is not None else '未知'}"
                ),
                str(feature_set.query.limit),
                "是（数量为下界）" if source.is_truncated else "否",
                (
                    str(source.dataset_record_count)
                    if source.dataset_record_count is not None
                    else "未标注"
                ),
                "合成" if source.is_synthetic else "外部/正式",
                fallback,
            ]
        )
    _add_table(
        document,
        [
            "分组",
            "Provider",
            "数据集",
            "版本",
            "查询时间",
            "数据更新时间",
            "CRS",
            "返回/可用",
            "查询上限",
            "截断",
            "数据集总量",
            "性质",
            "降级来源",
        ],
        source_rows,
        [600, 550, 1000, 750, 900, 900, 600, 600, 450, 450, 550, 500, 1510],
    )
    quality_notices = {
        feature_set.source.quality_notice
        for feature_set in result.poi_evidence.feature_sets
        if feature_set.source.quality_notice is not None
    }
    if quality_notices:
        _add_table_note(
            document,
            "数据质量边界: " + "；".join(sorted(quality_notices)),
        )
    if any(
        feature_set.source.is_truncated
        for feature_set in result.poi_evidence.feature_sets
    ):
        _add_table_note(
            document,
            "截断说明: 返回数量、密度和评分输入仅代表查询上限内的结果下界，需扩大分页或补充正式数据后复核。",
        )

    if result.site_score_report is not None:
        document.add_heading("3.3 GIS+POI 组合评分", level=2)
        site_score = result.site_score_report
        _add_label_detail_table(
            document,
            [
                ("组合评分版本", site_score.scoring_version),
                ("GIS 组件", f"{site_score.gis_component.score:.2f}"),
                ("POI 组件", f"{site_score.poi_score:.2f}"),
                (
                    "组件权重",
                    f"GIS={site_score.gis_weight:.3f}, POI={site_score.poi_weight:.3f}",
                ),
                ("组合总分", f"{site_score.total_score:.2f}"),
            ],
        )

    document.add_heading("3.4 政策规则证据", level=2)
    findings = result.policy_evidence.rule_findings
    if findings:
        policy_rows = [
            [
                f"{finding.rule_id}@{finding.rule_version}",
                finding.outcome.value,
                finding.policy_clause,
                finding.source_uri,
            ]
            for finding in findings
        ]
        _add_table(
            document,
            ["规则版本", "结果", "政策条款", "来源"],
            policy_rows,
            [2400, 1700, 2360, 2900],
        )
    else:
        paragraph = document.add_paragraph(
            "当前已评估规则未命中；该事实不等于整体合规。"
        )
        paragraph.paragraph_format.space_after = Pt(6)


def _add_review_section(document: DocumentType, state: AgentState) -> None:
    document.add_heading("4. 证据审查事项", level=1)
    rows = [
        [
            issue.severity.value,
            issue.issue_code,
            issue.parcel_id or "全局",
            issue.message,
        ]
        for issue in state.evidence_review_report.issues
    ]
    _add_table(
        document,
        ["级别", "代码", "地块", "说明"],
        rows,
        [1200, 2500, 1400, 4260],
    )


def _add_source_section(document: DocumentType, state: AgentState) -> None:
    document.add_heading("5. 来源与复核边界", level=1)
    sources = set()
    for result in state.results:
        sources.update(f"GIS:{item}" for item in result.gis_evidence.dataset_ids)
        sources.update(
            (
                f"POI:{feature_set.source.provider.value}:"
                f"{feature_set.source.dataset_id}:"
                f"{feature_set.source.dataset_version or 'unversioned'}:"
                f"queried_at={feature_set.source.queried_at.isoformat()}:"
                "dataset_updated_at="
                f"{feature_set.source.dataset_updated_at.isoformat() if feature_set.source.dataset_updated_at else 'unknown'}:"
                f"returned={feature_set.source.record_count}:"
                f"available={feature_set.source.available_record_count if feature_set.source.available_record_count is not None else 'unknown'}:"
                f"truncated={str(feature_set.source.is_truncated).lower()}"
            )
            for feature_set in result.poi_evidence.feature_sets
        )
        sources.update(
            f"POLICY:{finding.policy_id}:{finding.policy_clause}:{finding.source_uri}"
            for finding in result.policy_evidence.rule_findings
        )
    for source in sorted(sources):
        paragraph = document.add_paragraph(source)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(3)
    final = document.add_paragraph()
    final.paragraph_format.space_before = Pt(8)
    _add_run(
        final,
        "最终判断必须由具备权限的人员结合完整政策、现状资料和法定程序复核。",
        10.5,
        RGBColor(0x9B, 0x1C, 0x1C),
        bold=True,
    )


def _add_label_detail_table(
    document: DocumentType,
    rows: list[tuple[str, str]],
) -> None:
    _add_table(
        document,
        ["字段", "内容"],
        [[label, value] for label, value in rows],
        [2700, 6660],
    )


def _add_table(
    document: DocumentType,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
) -> None:
    if sum(widths_dxa) != _CONTENT_WIDTH_DXA:
        raise ValueError("Word 表格列宽之和必须为 9360 DXA")
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
        _shade_cell(cell, "F2F4F7")
        for run in cell.paragraphs[0].runs:
            _format_run(run, "Arial", 10, RGBColor(0, 0, 0), bold=True)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = str(text)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.10
                for run in paragraph.runs:
                    _format_run(run, "Arial", 9.5, RGBColor(0, 0, 0))
    _set_table_geometry(table, widths_dxa)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    table_properties = table._tbl.tblPr
    _set_child_value(table_properties, "tblW", "w", _CONTENT_WIDTH_DXA)
    _set_child_value(table_properties, "tblInd", "w", _TABLE_INDENT_DXA)
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            _set_child_value(properties, "tcW", "w", width)
            margins = properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                properties.append(margins)
            for side, value in _CELL_MARGIN_DXA.items():
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def _set_child_value(parent, tag: str, attribute: str, value: int) -> None:
    child = parent.find(qn(f"w:{tag}"))
    if child is None:
        child = OxmlElement(f"w:{tag}")
        parent.append(child)
    child.set(qn(f"w:{attribute}"), str(value))
    child.set(qn("w:type"), "dxa")


def _add_table_note(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    _add_run(paragraph, text, 9, _GRAY)


def _set_style_font(style, name, size, color, *, bold=False) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def _add_run(paragraph, text, size, color, *, bold=False):
    run = paragraph.add_run(text)
    _format_run(run, "Arial", size, color, bold=bold)
    return run


def _format_run(run, name, size, color, *, bold=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _format_optional_score(value: float | None) -> str:
    return "未生成" if value is None else f"{value:.2f}"
